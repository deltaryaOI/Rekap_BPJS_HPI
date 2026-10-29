import streamlit as st
import pandas as pd
import pdfplumber
import json
import os
import re
import io
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docxtpl import DocxTemplate
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Master Rekap & Generator Surat BPJS", layout="wide")
st.title("📄 Master Rekap & Generator Dokumen BPJS")

# ==============================================================================
# FUNGSI KEAMANAN & HELPER
# ==============================================================================
def safe_int(val):
    try:
        if pd.isna(val) or val is None or str(val).strip() == "" or str(val).strip().lower() == "none":
            return 0
        return int(float(val))
    except (ValueError, TypeError):
        return 0

# Fungsi Ekstraksi Periode ("19 Oktober 2026" -> "Oktober 2026")
def extract_periode(tanggal_str):
    if not tanggal_str:
        return "-"
    parts = tanggal_str.strip().split()
    if len(parts) >= 3:
        return f"{parts[-2]} {parts[-1]}"
    elif len(parts) == 2:
        return tanggal_str
    return tanggal_str

# Fungsi Terbilang Rupiah
def terbilang(n):
    satuan = ["", "Satu", "Dua", "Tiga", "Empat", "Lima", "Enam", "Tujuh", "Delapan", "Sembilan", "Sepuluh", "Sebelas"]
    n = int(abs(n))
    if n < 12:
        return satuan[n]
    elif n < 20:
        return terbilang(n - 10) + " Belas"
    elif n < 100:
        return terbilang(n // 10) + " Puluh " + terbilang(n % 10)
    elif n < 200:
        return "Seratus " + terbilang(n - 100)
    elif n < 1000:
        return terbilang(n // 100) + " Ratus " + terbilang(n % 100)
    elif n < 2000:
        return "Seribu " + terbilang(n - 1000)
    elif n < 1000000:
        return terbilang(n // 1000) + " Ribu " + terbilang(n % 1000)
    elif n < 1000000000:
        return terbilang(n // 1000000) + " Juta " + terbilang(n % 1000000)
    elif n < 1000000000000:
        return terbilang(n // 1000000000) + " Miliar " + terbilang(n % 1000000000)
    return str(n)

# Fungsi Border Tabel Word
def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tcBorders>
    ''' % nsdecls('w'))
    tcPr.append(tcBorders)

# ==============================================================================
# DATABASE KONTRAK & SIDEBAR SURAT WORD
# ==============================================================================
@st.cache_data
def load_database():
    possible_paths = [
        "database_kontrak.json",
        os.path.join("project-bpjs", "database_kontrak.json")
    ]
    for path in possible_paths:
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
    return {}

db_kontrak = load_database()

st.sidebar.header("📌 Pengaturan Dokumen Word")
wilayah_list = list(db_kontrak.keys()) if db_kontrak else ["CAB. PALEMBANG", "CAB. MEDAN", "CAB. SIBOLGA NIAS", "CAB. BABEL"]
selected_wilayah = st.sidebar.selectbox("Pilih Cabang / Wilayah", wilayah_list)

wilayah_display = re.sub(r'(?i)\bcab\.?\s*', '', selected_wilayah).strip()

kontrak_options = db_kontrak.get(selected_wilayah, [])
if not kontrak_options:
    kontrak_options = ["-"]

st.sidebar.markdown("---")
st.sidebar.header("📝 Informasi Header Surat")
nomor_surat = st.sidebar.text_input("Nomor Surat", "001/SP/HPI/X/2026")
tanggal_surat = st.sidebar.text_input("Tanggal Surat", "19 Oktober 2026")
nama_manager = st.sidebar.text_input("Nama Manager Cabang", "NAMA MANAGER")

tab1, tab2 = st.tabs(["🏢 BPJS Ketenagakerjaan (EPS & NPP)", "🏥 BPJS Kesehatan (Billing Statement)"])

# ==============================================================================
# TAB 1: BPJS KETENAGAKERJAAN
# ==============================================================================
with tab1:
    st.subheader("Rekap PDF & Generator Surat BPJS Ketenagakerjaan")
    uploaded_files_tk = st.file_uploader(
        "Upload File PDF BPJS TK (Bisa banyak file sekaligus)", 
        type=["pdf"], 
        accept_multiple_files=True,
        key="uploader_tk"
    )

    if uploaded_files_tk:
        st.info(f"Memproses {len(uploaded_files_tk)} file BPJS Ketenagakerjaan...")
        all_data_tk = []
        
        def parse_amount(val_str):
            if not val_str:
                return 0.0
            clean_str = val_str.strip()
            if '.' in clean_str and ',' in clean_str:
                if clean_str.rfind(',') > clean_str.rfind('.'):
                    clean_str = clean_str.replace('.', '').replace(',', '.')
                else:
                    clean_str = clean_str.replace(',', '')
            elif ',' in clean_str and '.' not in clean_str:
                clean_str = clean_str.replace('.', '').replace(',', '.')
            else:
                clean_str = clean_str.replace(',', '')
            try:
                return float(clean_str)
            except ValueError:
                return 0.0

        for file in uploaded_files_tk:
            try:
                with pdfplumber.open(file) as pdf:
                    raw_text = ""
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            raw_text += text + "\n"

                clean_text = raw_text.replace('\n', ' ').replace('|', '')
                default_spk = kontrak_options[0] if (kontrak_options and len(kontrak_options) > 0) else "-"

                data_tk = {
                    "Nama Perusahaan": "-",
                    "NPP": "-",
                    "Nomor Kontrak": str(default_spk),
                    "Kode Iuran": "-",
                    "Jumlah Tenaga Kerja": 0, 
                    "Total Iuran JKK": 0.0,
                    "Total Iuran JHT": 0.0,
                    "Total Iuran JKM": 0.0,
                    "Total Iuran JPN": 0.0,
                    "Total Pembayaran": 0.0
                }

                # 1. Nama Perusahaan
                match_nama_sipp = re.search(r'Nama.?\s:\s*(.*?)\s(?:Total|Jumlah)\s*Tenaga', clean_text, re.IGNORECASE)
                match_nama_eps = re.search(r'PEMBAYARAN\s*UNTUK\s+(.*?)\s+NPP', clean_text, re.IGNORECASE)
                match_nama_eps_ringkas = re.search(r'Nama\s+Perusahaan\s*:\s*(.*?)(?:\s+Total\s+Tenaga|\s+Bulan|$)', clean_text, re.IGNORECASE)

                if match_nama_sipp:
                    nama_pt = match_nama_sipp.group(1).replace('Perusahaan', '').replace('Kerja', '').strip()
                    data_tk["Nama Perusahaan"] = re.sub(r'\s+', ' ', nama_pt)
                elif match_nama_eps:
                    data_tk["Nama Perusahaan"] = re.sub(r'\s+', ' ', match_nama_eps.group(1).strip())
                elif match_nama_eps_ringkas:
                    data_tk["Nama Perusahaan"] = re.sub(r'\s+', ' ', match_nama_eps_ringkas.group(1).strip())

                # 2. NPP
                match_npp = re.search(r'NPP\s*:\s*(\d+)', clean_text, re.IGNORECASE)
                if match_npp:
                    data_tk["NPP"] = str(match_npp.group(1).strip())

                # 3. Kode Iuran
                match_kode = re.search(r'(?:KODE\s*IURAN\s*:?|KODE\s*PEMBAYARAN)\s*([0-9\s]{10,20})', clean_text, re.IGNORECASE)
                if match_kode: 
                    data_tk["Kode Iuran"] = match_kode.group(1).replace(' ', '').strip()

                # 4. Jumlah Tenaga Kerja
                match_tk = re.search(r'(?:JUMLAH|Total)\s*TENAGA\s*KERJA\s*:\s*(\d+)', clean_text, re.IGNORECASE)
                if match_tk:
                    data_tk["Jumlah Tenaga Kerja"] = int(match_tk.group(1).strip())

                # 5. Rincian Iuran
                match_jkk = re.search(r'JKK.*?(?:Rp\.?|Rp)?\s*([\d,.]+)', clean_text, re.IGNORECASE)
                if match_jkk:
                    data_tk["Total Iuran JKK"] = parse_amount(match_jkk.group(1))

                match_jht = re.search(r'JHT.*?(?:Rp\.?|Rp)?\s*([\d,.]+)', clean_text, re.IGNORECASE)
                if match_jht:
                    data_tk["Total Iuran JHT"] = parse_amount(match_jht.group(1))

                match_jkm = re.search(r'JKM.*?(?:Rp\.?|Rp)?\s*([\d,.]+)', clean_text, re.IGNORECASE)
                if match_jkm:
                    data_tk["Total Iuran JKM"] = parse_amount(match_jkm.group(1))

                match_jpn = re.search(r'(?:JPN|JP)\s.*?(?:Rp\.?|Rp)?\s*([\d,.]+)', clean_text, re.IGNORECASE)
                if match_jpn:
                    data_tk["Total Iuran JPN"] = parse_amount(match_jpn.group(1))

                # 6. Total Pembayaran
                match_total = re.search(r'(?:TOTAL\s*YANG\s*DIBAYARKAN|TOTAL\s*IURAN)\s*(?:\(\s*PEMBULATAN\s*\))?\s*:?\s*(?:Rp\.?|Rp)?\s*([\d,.]+)', clean_text, re.IGNORECASE)
                if match_total:
                    data_tk["Total Pembayaran"] = parse_amount(match_total.group(1))

                all_data_tk.append(data_tk)
                
            except Exception as e:
                st.error(f"Terjadi kesalahan saat membaca file {file.name}: {e}")

        if all_data_tk:
            df_raw = pd.DataFrame(all_data_tk)
            
            st.subheader("✏️ Edit Data & Nomor Kontrak (SPK)")
            st.caption("💡 Klik 2x sel di bawah ini untuk mengedit data secara langsung.")

            with st.expander("📌 Lihat Daftar Nomor Kontrak (SPK) Cabang Ini"):
                for opt in kontrak_options:
                    st.code(opt, language="text")

            column_config = {
                "Nama Perusahaan": st.column_config.TextColumn("Nama Perusahaan"),
                "NPP": st.column_config.TextColumn("NPP"),
                "Nomor Kontrak": st.column_config.TextColumn("Nomor Kontrak (SPK)", width="large"),
                "Kode Iuran": st.column_config.TextColumn("Kode Iuran"),
                "Jumlah Tenaga Kerja": st.column_config.NumberColumn("Jml TK", format="%d"),
                "Total Iuran JKK": st.column_config.NumberColumn("Total JKK", format="Rp %.2f"),
                "Total Iuran JHT": st.column_config.NumberColumn("Total JHT", format="Rp %.2f"),
                "Total Iuran JKM": st.column_config.NumberColumn("Total JKM", format="Rp %.2f"),
                "Total Iuran JPN": st.column_config.NumberColumn("Total JPN", format="Rp %.2f"),
                "Total Pembayaran": st.column_config.NumberColumn("Total Pembayaran (Rp)", format="Rp %.2f")
            }

            edited_df = st.data_editor(
                df_raw,
                column_config=column_config,
                num_rows="dynamic",
                use_container_width=True,
                key=f"editor_combine_{selected_wilayah}_{len(uploaded_files_tk)}"
            )

            edited_df.reset_index(drop=True, inplace=True)
            edited_df["NO"] = [idx + 1 for idx in range(len(edited_df))]

            total_iuran_sum = float(edited_df["Total Pembayaran"].sum())
            total_tk_sum = int(edited_df["Jumlah Tenaga Kerja"].sum())
            terbilang_iuran = (terbilang(total_iuran_sum).strip() + " Rupiah") if total_iuran_sum > 0 else "-"

            st.markdown(f"### **Total Anggaran Iuran:** Rp {total_iuran_sum:,.0f}".replace(",", "."))
            st.markdown(f"**Terbilang:** *{terbilang_iuran}*")

            st.subheader("📥 Unduh Hasil Dokumen")
            col1, col2 = st.columns(2)

            # 1. Download Excel
            with col1:
                st.markdown("##### 📊 Export Master Excel")
                clean_filename_wilayah = wilayah_display.replace(".", "").replace(" ", "_")
                excel_filename = f"Rekap_BPJS_TK_{clean_filename_wilayah}.xlsx"
                
                output_excel = io.BytesIO()
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Rekap BPJS TK"

                font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                font_data = Font(name="Calibri", size=11)
                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'),
                    right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'),
                    bottom=Side(style='thin', color='D9D9D9')
                )

                headers = ["NO", "Nama Perusahaan", "NPP", "Nomor Kontrak", "Kode Iuran", "Jml TK", "Total JKK", "Total JHT", "Total JKM", "Total JPN", "Total Pembayaran"]
                ws.append(headers)
                for col_num in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                records = edited_df.to_dict(orient="records")
                for r_idx, row in enumerate(records, start=2):
                    ws.cell(row=r_idx, column=1, value=int(row["NO"])).alignment = Alignment(horizontal="center")
                    ws.cell(row=r_idx, column=2, value=str(row["Nama Perusahaan"])).alignment = Alignment(horizontal="left")
                    ws.cell(row=r_idx, column=3, value=str(row["NPP"])).alignment = Alignment(horizontal="center")
                    ws.cell(row=r_idx, column=4, value=str(row["Nomor Kontrak"])).alignment = Alignment(horizontal="left")
                    ws.cell(row=r_idx, column=5, value=str(row["Kode Iuran"])).alignment = Alignment(horizontal="center")
                    ws.cell(row=r_idx, column=6, value=safe_int(row["Jumlah Tenaga Kerja"])).alignment = Alignment(horizontal="center")
                    
                    for c_off, col_name in enumerate(["Total Iuran JKK", "Total Iuran JHT", "Total Iuran JKM", "Total Iuran JPN", "Total Pembayaran"], start=7):
                        c_cell = ws.cell(row=r_idx, column=c_off, value=float(row[col_name]))
                        c_cell.number_format = '"Rp "#,##0.00'
                        c_cell.alignment = Alignment(horizontal="right")

                    for c_idx in range(1, 12):
                        c = ws.cell(row=r_idx, column=c_idx)
                        c.font = font_data
                        c.border = thin_border

                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

                wb.save(output_excel)
                output_excel.seek(0)

                st.download_button(
                    label="📥 Download Master Excel (.xlsx)",
                    data=output_excel.getvalue(),
                    file_name=excel_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            # 2. Download Word Docx
            with col2:
                st.markdown("##### 📝 Export Surat Word (.docx)")
                template_path = os.path.join("templates", "Template_DEFAULT_ketenagakerjaan.docx")
                if not os.path.exists(template_path):
                    template_path = os.path.join("templates", "Template_DEFAULT.docx")

                if os.path.exists(template_path):
                    doc_tpl = DocxTemplate(template_path)
                    periode_surat = extract_periode(tanggal_surat)

                    context = {
                        "Wilayah": wilayah_display,
                        "Tanggal": tanggal_surat,
                        "Periode": periode_surat,
                        "NomorSurat": nomor_surat,
                        "Terbilang": terbilang_iuran,
                        "NamaManager": nama_manager,
                        "TOTAL_TK": str(total_tk_sum),
                        "TOTAL_IURAN": f"Rp {total_iuran_sum:,.0f}".replace(",", ".")
                    }
                    doc_tpl.render(context)
                    
                    temp_stream = io.BytesIO()
                    doc_tpl.save(temp_stream)
                    temp_stream.seek(0)

                    doc_native = docx.Document(temp_stream)

                    target_p = None
                    for p in doc_native.paragraphs:
                        if "sebesar" in p.text.lower() or "anggaran" in p.text.lower():
                            target_p = p
                            break

                    for tbl in list(doc_native.tables):
                        tbl._element.getparent().remove(tbl._element)

                    if target_p:
                        spacer_top = doc_native.add_paragraph()
                        spacer_top.paragraph_format.space_before = Pt(4)
                        spacer_top.paragraph_format.space_after = Pt(4)
                        target_p._p.addnext(spacer_top._p)
                        insert_point = spacer_top
                    else:
                        insert_point = target_p

                    new_table = doc_native.add_table(rows=0, cols=5)
                    new_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    if insert_point:
                        insert_point._p.addnext(new_table._tbl)

                    hdr_cells = new_table.add_row().cells
                    headers_docx = ["NO", "NPP", "Nomor Kontrak", "Jml TK", "Total Iuran"]
                    widths = [Inches(0.5), Inches(1.3), Inches(3.1), Inches(0.8), Inches(1.8)]
                    
                    for idx, text in enumerate(headers_docx):
                        hdr_cells[idx].text = text
                        hdr_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
                        set_cell_border(hdr_cells[idx])
                        hdr_cells[idx].width = widths[idx]

                    records = edited_df.to_dict(orient="records")
                    for row_data in records:
                        val_iuran = float(row_data["Total Pembayaran"])
                        formatted_iuran = f"Rp\t{val_iuran:,.0f}".replace(",", ".")

                        row_cells = new_table.add_row().cells
                        row_cells[0].text = str(row_data["NO"])
                        row_cells[1].text = str(row_data["NPP"])
                        row_cells[2].text = str(row_data["Nomor Kontrak"])
                        row_cells[3].text = str(safe_int(row_data["Jumlah Tenaga Kerja"]))
                        
                        p_iuran = row_cells[4].paragraphs[0]
                        p_iuran.text = formatted_iuran
                        p_iuran.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p_iuran.paragraph_format.tab_stops.add_tab_stop(Inches(1.75), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)

                        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                        row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                        for c_idx in range(5):
                            set_cell_border(row_cells[c_idx])
                            row_cells[c_idx].width = widths[c_idx]

                    total_cells = new_table.add_row().cells
                    total_cells[0].text = "Jumlah"
                    total_cells[3].text = str(total_tk_sum)

                    p_tot_iuran = total_cells[4].paragraphs[0]
                    p_tot_iuran.text = f"Rp\t{total_iuran_sum:,.0f}".replace(",", ".")
                    p_tot_iuran.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_tot_iuran.paragraph_format.tab_stops.add_tab_stop(Inches(1.75), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
                    if p_tot_iuran.runs:
                        p_tot_iuran.runs[0].font.bold = True

                    total_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    total_cells[0].paragraphs[0].runs[0].font.bold = True
                    total_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    total_cells[3].paragraphs[0].runs[0].font.bold = True

                    for c_idx in range(5):
                        set_cell_border(total_cells[c_idx])
                        total_cells[c_idx].width = widths[c_idx]

                    spacer_bottom = doc_native.add_paragraph()
                    spacer_bottom.paragraph_format.space_before = Pt(6)
                    spacer_bottom.paragraph_format.space_after = Pt(6)
                    new_table._tbl.addnext(spacer_bottom._p)

                    doc_io = io.BytesIO()
                    doc_native.save(doc_io)
                    doc_io.seek(0)
                    
                    word_filename = f"Surat_BPJS_TK_{clean_filename_wilayah}.docx"
                    st.download_button(
                        label="📥 Download Surat Word (.docx)",
                        data=doc_io.getvalue(),
                        file_name=word_filename,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                else:
                    st.error("Template Word Ketenagakerjaan belum ditemukan di `templates/`!")

# ==============================================================================
# TAB 2: BPJS KESEHATAN
# ==============================================================================
with tab2:
    st.subheader("Rekap PDF Billing Statement & Generator Surat BPJS Kesehatan")
    uploaded_files_kes = st.file_uploader(
        "Upload File PDF Billing Statement BPJS Kesehatan (Bisa banyak file sekaligus)", 
        type=["pdf"], 
        accept_multiple_files=True, 
        key="uploader_kes"
    )

    if uploaded_files_kes:
        st.info(f"Memproses {len(uploaded_files_kes)} file BPJS Kesehatan...")
        all_data_kes = []

        for file in uploaded_files_kes:
            try:
                with pdfplumber.open(file) as pdf:
                    raw_text = ""
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            raw_text += text + "\n"

                clean_text = raw_text.replace('\n', ' ').replace('|', '')

                data_kes = {
                    "Nama File": file.name,
                    "Nama Perusahaan": "-",
                    "Nomor VA": "-",
                    "Jumlah Tenaga Kerja": 0,
                    "Total Iuran": 0.0
                }

                # 1. NAMA PERUSAHAAN
                match_nama = re.search(r'Yth\.\s*(.*?)(?:\n\s*JL\.|\s*JL\.|\s*Nomor Virtual|\s*Pelaporan|\n\n|$)', raw_text, re.IGNORECASE | re.DOTALL)
                if match_nama:
                    nama_bu = match_nama.group(1).replace('\n', ' ').strip()
                    nama_bu = re.sub(r'(?:Surat Tagihan|Billing Statement|Pelaporan Penagihan|JKN-KIS).*$', '', nama_bu, flags=re.IGNORECASE).strip()
                    data_kes["Nama Perusahaan"] = re.sub(r'\s+', ' ', nama_bu)

                # 2. NOMOR VA
                match_va = re.search(r'(?:BNI|BRI|MDR|BTN|Virtual Account)\s*:?\s*([0-9]{13,16})', clean_text, re.IGNORECASE)
                if match_va:
                    va_clean = match_va.group(1).strip()
                    data_kes["Nomor VA"] = va_clean

                # 3. JUMLAH TENAGA KERJA
                match_pegawai = re.search(r'(?:Rp\.?|Rp)?\s*[\d.]+\s*(?:Rp\.?|Rp)?\s*(\d+)\s+\d+\s+\d{2}/\d{2}/\d{4}', clean_text, re.IGNORECASE)
                if match_pegawai:
                    data_kes["Jumlah Tenaga Kerja"] = int(match_pegawai.group(1).strip())
                else:
                    match_peg_alt = re.search(r'(?:Jumlah Pegawai|Number of Employees).*?\s+(\d+)', clean_text, re.IGNORECASE)
                    if match_peg_alt:
                        data_kes["Jumlah Tenaga Kerja"] = int(match_peg_alt.group(1).strip())

                # 4. TOTAL IURAN
                match_total = re.search(r'(?:TOTAL TAGIHAN|Total Balance).?(?:Rp\.?|Rp)\s([\d.]+)', clean_text, re.IGNORECASE)
                if not match_total:
                    match_total = re.search(r'(?:Premi/Iuran|Contribution).?(?:Rp\.?|Rp)\s([\d.]+)', clean_text, re.IGNORECASE)

                if match_total:
                    raw_angka = match_total.group(1).strip()
                    clean_angka = raw_angka.replace('.', '')
                    data_kes["Total Iuran"] = float(clean_angka)

            except Exception as e:
                st.error(f"Terjadi kesalahan saat membaca file {file.name}: {e}")

            all_data_kes.append(data_kes)

        if all_data_kes:
            df_kes = pd.DataFrame(all_data_kes)
            df_kes.reset_index(drop=True, inplace=True)
            df_kes["NO"] = [idx + 1 for idx in range(len(df_kes))]

            st.success("Ekstraksi BPJS Kesehatan Selesai!")
            st.write("### Preview Master Rekap BPJS Kesehatan:")

            def format_indo(angka):
                return f"{angka:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            def format_bulat_indo(angka):
                return f"{angka:,}".replace(",", ".")

            st.dataframe(df_kes.style.format({
                "Total Iuran": format_indo,
                "Jumlah Tenaga Kerja": format_bulat_indo
            }))

            total_iuran_kes = float(df_kes["Total Iuran"].sum())
            total_tk_kes = int(df_kes["Jumlah Tenaga Kerja"].sum())
            terbilang_kes = (terbilang(total_iuran_kes).strip() + " Rupiah") if total_iuran_kes > 0 else "-"

            st.markdown(f"### **Total Tagihan BPJS Kesehatan:** Rp {total_iuran_kes:,.0f}".replace(",", "."))
            st.markdown(f"**Terbilang:** *{terbilang_kes}*")

            st.subheader("📥 Unduh Hasil Dokumen BPJS Kesehatan")
            col1, col2 = st.columns(2)

            # 1. Excel BPJS Kesehatan
            with col1:
                st.markdown("##### 📊 Export Master Excel")
                clean_filename_wilayah = wilayah_display.replace(".", "").replace(" ", "_")
                excel_buffer_kes = io.BytesIO()
                df_kes[["NO", "Nama File", "Nama Perusahaan", "Nomor VA", "Jumlah Tenaga Kerja", "Total Iuran"]].to_excel(excel_buffer_kes, index=False, engine='openpyxl')

                st.download_button(
                    label="📥 Download Master Excel BPJS Kesehatan",
                    data=excel_buffer_kes.getvalue(),
                    file_name=f"Master_Rekap_BPJS_Kesehatan_{clean_filename_wilayah}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            # 2. Export Word Docx BPJS Kesehatan
            with col2:
                st.markdown("##### 📝 Export Surat Word (.docx)")
                template_kes_path = os.path.join("templates", "Template_DEFAULT_kesehatan.docx")
                if not os.path.exists(template_kes_path):
                    template_kes_path = os.path.join("templates", "Template_DEFAULT.docx")

                if os.path.exists(template_kes_path):
                    doc_tpl = DocxTemplate(template_kes_path)
                    periode_kes = extract_periode(tanggal_surat)

                    context = {
                        "Wilayah": wilayah_display,
                        "Tanggal": tanggal_surat,
                        "Periode": periode_kes,
                        "NomorSurat": nomor_surat,
                        "Terbilang": terbilang_kes,
                        "NamaManager": nama_manager,
                        "TOTAL_TK": str(total_tk_kes),
                        "TOTAL_IURAN": f"Rp {total_iuran_kes:,.0f}".replace(",", ".")
                    }
                    doc_tpl.render(context)
                    
                    temp_stream = io.BytesIO()
                    doc_tpl.save(temp_stream)
                    temp_stream.seek(0)

                    doc_native = docx.Document(temp_stream)
                    records_kes = df_kes.to_dict(orient="records")

                    target_p = None
                    for p in doc_native.paragraphs:
                        if "sebesar" in p.text.lower() or "anggaran" in p.text.lower():
                            target_p = p
                            break

                    for tbl in list(doc_native.tables):
                        tbl._element.getparent().remove(tbl._element)

                    if target_p:
                        spacer_top_kes = doc_native.add_paragraph()
                        spacer_top_kes.paragraph_format.space_before = Pt(4)
                        spacer_top_kes.paragraph_format.space_after = Pt(4)
                        target_p._p.addnext(spacer_top_kes._p)
                        insert_point_kes = spacer_top_kes
                    else:
                        insert_point_kes = target_p

                    new_table = doc_native.add_table(rows=0, cols=5)
                    new_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    if insert_point_kes:
                        insert_point_kes._p.addnext(new_table._tbl)

                    hdr_cells = new_table.add_row().cells
                    headers_kes_docx = ["NO", "NAMA BADAN USAHA/SPK", "NOMOR VA", "JML TK", "TOTAL TAGIHAN"]
                    widths_kes = [Inches(0.5), Inches(3.0), Inches(1.8), Inches(0.8), Inches(1.6)]
                    
                    for idx, text in enumerate(headers_kes_docx):
                        hdr_cells[idx].text = text
                        hdr_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
                        set_cell_border(hdr_cells[idx])
                        hdr_cells[idx].width = widths_kes[idx]

                    for r_data in records_kes:
                        val_tagihan = float(r_data["Total Iuran"])
                        formatted_tagihan = f"Rp\t{val_tagihan:,.0f}".replace(",", ".")

                        row_cells = new_table.add_row().cells
                        row_cells[0].text = str(r_data["NO"])
                        row_cells[1].text = str(r_data["Nama Perusahaan"])
                        row_cells[2].text = str(r_data["Nomor VA"])
                        row_cells[3].text = str(safe_int(r_data["Jumlah Tenaga Kerja"]))
                        
                        p_tagihan = row_cells[4].paragraphs[0]
                        p_tagihan.text = formatted_tagihan
                        p_tagihan.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p_tagihan.paragraph_format.tab_stops.add_tab_stop(Inches(1.55), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)

                        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                        for c_idx in range(5):
                            set_cell_border(row_cells[c_idx])
                            row_cells[c_idx].width = widths_kes[c_idx]

                    total_cells = new_table.add_row().cells
                    total_cells[0].text = "Jumlah"
                    total_cells[3].text = str(total_tk_kes)

                    p_tot_kes = total_cells[4].paragraphs[0]
                    p_tot_kes.text = f"Rp\t{total_iuran_kes:,.0f}".replace(",", ".")
                    p_tot_kes.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_tot_kes.paragraph_format.tab_stops.add_tab_stop(Inches(1.55), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
                    if p_tot_kes.runs:
                        p_tot_kes.runs[0].font.bold = True

                    total_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    total_cells[0].paragraphs[0].runs[0].font.bold = True
                    total_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    total_cells[3].paragraphs[0].runs[0].font.bold = True

                    for c_idx in range(5):
                        set_cell_border(total_cells[c_idx])
                        total_cells[c_idx].width = widths_kes[c_idx]

                    spacer_bottom_kes = doc_native.add_paragraph()
                    spacer_bottom_kes.paragraph_format.space_before = Pt(6)
                    spacer_bottom_kes.paragraph_format.space_after = Pt(6)
                    new_table._tbl.addnext(spacer_bottom_kes._p)

                    doc_io_kes = io.BytesIO()
                    doc_native.save(doc_io_kes)
                    doc_io_kes.seek(0)
                    
                    word_filename_kes = f"Surat_BPJS_Kesehatan_{clean_filename_wilayah}.docx"
                    st.download_button(
                        label="📥 Download Surat Word (.docx)",
                        data=doc_io_kes.getvalue(),
                        file_name=word_filename_kes,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                else:
                    st.error("Template Word belum ditemukan di `templates/`!")
